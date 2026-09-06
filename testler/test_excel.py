import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import os
import time
import database
import excel_manager
from openpyxl import load_workbook

def test_excel():
    print(">> EXCEL ENTEGRASYON TESTI BASLATILIYOR...")

    # 1. Dosyayı silerek başlayalım
    if os.path.exists(excel_manager.EXCEL_DOSYA_ADI):
        try:
            os.remove(excel_manager.EXCEL_DOSYA_ADI)
            print(f"[OK] Eski {excel_manager.EXCEL_DOSYA_ADI} silindi.")
        except Exception as e:
            print(f"Eski dosya silinemedi: {e}")

    # 2. Excel dosyasını otomatik oluşturma testi
    excel_manager.excel_dosyasini_hazirla()
    assert os.path.exists(excel_manager.EXCEL_DOSYA_ADI), "Excel dosyası oluşturulamadı!"
    print("[OK] Dosya bulunamadığında otomatik oluşturma başarılı.")

    # 3. Araç girişi yap ve Excel'e yazıldığını doğrula
    giris_kaydi, _ = database.arac_giris_yap("34ABC29", time.time() - 25)
    print(f"[OK] 34ABC29 giriş yaptı (ID: {giris_kaydi['id']}).")

    wb = load_workbook(excel_manager.EXCEL_DOSYA_ADI)
    ws = wb.active
    assert ws.max_row >= 2, "Excel'e giriş satırı eklenmedi!"
    son_plaka = ws.cell(row=ws.max_row, column=2).value
    son_durum = ws.cell(row=ws.max_row, column=9).value
    assert son_plaka == "34ABC29"
    assert son_durum == "İçeride"
    print(f"[OK] Excel giriş kaydı doğrulandı: Plaka={son_plaka}, Durum={son_durum}")

    # 4. Araç çıkış yap ve Excel satırının güncellendiğini doğrula
    basarili, mesaj, cikis_kaydi = database.arac_cikis_yap("34ABC29", zorla=True)
    assert basarili == True
    print(f"[OK] 34ABC29 çıkış yaptı. Süre: {cikis_kaydi['toplam_sure_sn']:.1f} sn, Ücret: {cikis_kaydi['ucret']} TL")

    wb = load_workbook(excel_manager.EXCEL_DOSYA_ADI)
    ws = wb.active
    row_found = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == "34ABC29":
            row_found = r
            break
    assert row_found is not None
    durum = ws.cell(row=row_found, column=9).value
    ucret = ws.cell(row=row_found, column=8).value
    cikis_zamani = ws.cell(row=row_found, column=5).value
    assert durum == "Çıkış Yaptı"
    assert ucret > 0
    assert cikis_zamani != "-"
    print(f"[OK] Excel çıkış güncellemesi doğrulandı: Çıkış={cikis_zamani}, Ücret={ucret} TL, Durum={durum}")

    # 5. DOSYAYI KULLANICI YANLIŞLIKLA SİLERSE TESTİ!
    print("\n>> KULLANICI DOSYAYI YANLISLIKLA SILDI SENARYOSU:")
    os.remove(excel_manager.EXCEL_DOSYA_ADI)
    assert not os.path.exists(excel_manager.EXCEL_DOSYA_ADI)
    print("[OK] Dosya diskten silindi.")

    # Otomatik veya senkronize çağrısı
    toplam = database.excel_senkronize_et()
    assert os.path.exists(excel_manager.EXCEL_DOSYA_ADI), "Silinen dosya yeniden oluşturulamadı!"
    wb = load_workbook(excel_manager.EXCEL_DOSYA_ADI)
    ws = wb.active
    assert ws.max_row >= 2
    print(f"[OK] Silinen dosya tüm veritabanı kayıtlarıyla ({toplam} adet) sıfırdan ve eksiksiz geri oluşturuldu!")

    print("\n>>> TÜM EXCEL TESTLERİ EKSİKSİZ BAŞARIYLA GEÇTİ! <<<")

if __name__ == "__main__":
    test_excel()
